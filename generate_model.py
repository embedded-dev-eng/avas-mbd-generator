"""
generate_model.py
------------------
Outil d'automatisation Model-Based Design (MBD).

Lit :
  - data/functions.xlsx   -> architecture fonctionnelle (blocs, E/S, enchainement)
  - data/parameters.xml   -> parametres de chaque fonction (equivalent "data dictionary")

Genere :
  - output/build_avas_model.m  -> script MATLAB qui construit automatiquement
                                   le modele Simulink (blocs + interconnexions + params)
  - output/architecture_preview.svg -> apercu visuel du modele (sans besoin de MATLAB)

Usage :
  python3 generate_model.py
"""

import openpyxl
import xml.etree.ElementTree as ET
from pathlib import Path
import graphviz

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)


def read_functions(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Functions"]
    functions = []
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        if not record.get("FunctionID"):
            continue
        record["Inputs"] = [s.strip() for s in str(record["Inputs"]).split(",")]
        record["Outputs"] = [s.strip() for s in str(record["Outputs"]).split(",")]
        functions.append(record)
    return functions


def read_parameters(xml_path):
    tree = ET.parse(xml_path)
    root = tree.getroot()
    params_by_function = {}
    for func_el in root.findall("Function"):
        fid = func_el.get("id")
        params = []
        for p in func_el.findall("Parameter"):
            params.append({
                "name": p.get("name"),
                "value": p.get("value"),
                "type": p.get("type", "double"),
            })
        params_by_function[fid] = params
    return params_by_function


def generate_matlab_script(functions, params_by_function, out_path):
    """Genere un script .m qui construit le modele Simulink automatiquement,
    bloc par bloc, avec ses ports et ses parametres, a partir des donnees lues."""

    model_name = "AVAS_ElectricVehicle"
    lines = []
    lines.append(f"% Script genere automatiquement par generate_model.py")
    lines.append(f"% Ne pas editer a la main : modifier data/functions.xlsx et data/parameters.xml")
    lines.append(f"% puis relancer generate_model.py")
    lines.append("")
    lines.append(f"modelName = '{model_name}';")
    lines.append("if bdIsLoaded(modelName)")
    lines.append("    close_system(modelName, 0);")
    lines.append("end")
    lines.append("new_system(modelName);")
    lines.append("open_system(modelName);")
    lines.append("")
    lines.append("%% --- Data dictionary : parametres issus de parameters.xml ---")
    for fid, params in params_by_function.items():
        for p in params:
            varname = f"{fid}_{p['name']}"
            lines.append(f"assignin('base', '{varname}', {p['value']});")
    lines.append("")

    lines.append("%% --- Creation des blocs Subsystem (un par fonction) ---")
    x = 30
    y = 30
    step_x = 220
    block_positions = {}
    for i, func in enumerate(functions):
        fid = func["FunctionID"]
        fname = func["FunctionName"]
        pos = f"[{x} {y} {x+140} {y+80}]"
        block_positions[fid] = (x, y)
        lines.append(f"add_block('simulink/Ports & Subsystems/Subsystem', "
                      f"[modelName '/{fname}'], 'Position', '{pos}');")
        # supprime les ports In1/Out1 par defaut livres avec le bloc Subsystem,
        # pour ne garder que les ports nommes definis dans functions.xlsx
        lines.append(f"delete_block([modelName '/{fname}/In1']);")
        lines.append(f"delete_block([modelName '/{fname}/Out1']);")
        # supprime aussi le fil par defaut qui reliait ces deux ports (sinon il
        # reste orphelin, affiche en rouge pointille dans Simulink)
        lines.append(f"try; delete_line(find_system([modelName '/{fname}'], "
                      f"'FindAll', 'on', 'SearchDepth', 1, 'Type', 'line')); catch; end")
        # ajoute les ports d'entree/sortie a l'interieur du sous-systeme
        for j, inp in enumerate(func["Inputs"]):
            lines.append(f"add_block('simulink/Sources/In1', "
                          f"[modelName '/{fname}/{inp}'], 'Position', '[30 {30+j*60} 60 {50+j*60}]');")
        for j, outp in enumerate(func["Outputs"]):
            lines.append(f"add_block('simulink/Sinks/Out1', "
                          f"[modelName '/{fname}/{outp}'], 'Position', '[300 {30+j*60} 330 {50+j*60}]');")
        x += step_x

    lines.append("")
    lines.append("%% --- Logique interne : SoundGenerator (F02) ---")
    lines.append("% Lookup Table 1-D : convertit VehicleSpeed_kmh en frequence audio (Hz),")
    lines.append("% pilotee par les seuils issus de parameters.xml (aucune valeur en dur).")
    lines.append("add_block('simulink/Lookup Tables/Lookup Table', "
                  "[modelName '/SoundGenerator/FrequencyMapping'], "
                  "'Position', '[150 65 230 95]');")
    lines.append("set_param([modelName '/SoundGenerator/FrequencyMapping'], "
                  "'InputValues', '[F02_ActivationSpeed_kmh_min F02_ActivationSpeed_kmh_max]', "
                  "'Table', '[F02_FrequencyBand_Hz_min F02_FrequencyBand_Hz_max]');")
    lines.append("add_line([modelName '/SoundGenerator'], 'VehicleSpeed_kmh/1', "
                  "'FrequencyMapping/1', 'autorouting', 'on');")
    lines.append("add_line([modelName '/SoundGenerator'], 'FrequencyMapping/1', "
                  "'AudioSignal_raw/1', 'autorouting', 'on');")

    lines.append("")
    lines.append("%% --- Interconnexion des blocs selon NextFunctionID ---")
    fname_by_id = {f["FunctionID"]: f["FunctionName"] for f in functions}
    for func in functions:
        nxt = func.get("NextFunctionID")
        if nxt and str(nxt).strip():
            src = fname_by_id[func["FunctionID"]]
            dst = fname_by_id[nxt]
            lines.append(f"add_line(modelName, '{src}/1', '{dst}/1', 'autorouting', 'on');")

    lines.append("")
    lines.append("Simulink.BlockDiagram.arrangeSystem(modelName);")
    lines.append(f"save_system(modelName, fullfile(fileparts(mfilename('fullpath')), '{model_name}.slx'));")
    lines.append(f"disp('Modele {model_name}.slx genere avec succes.');")

    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"OK: {out_path}")


def generate_preview_svg(functions, out_path):
    """Genere un apercu visuel du modele (utile pour le README GitHub,
    fonctionne sans MATLAB/Simulink installe)."""
    dot = graphviz.Digraph("architecture", format="svg")
    dot.attr(rankdir="LR", bgcolor="transparent", fontname="Helvetica")
    dot.attr("node", shape="box", style="rounded,filled", fillcolor="#f6f1e6",
             color="#2b3a55", fontname="Helvetica", fontsize="11", margin="0.25,0.15")
    dot.attr("edge", color="#6b6255", arrowsize="0.7")

    fname_by_id = {f["FunctionID"]: f["FunctionName"] for f in functions}
    for func in functions:
        label = f"{func['FunctionName']}\\n({func['FunctionID']})"
        dot.node(func["FunctionID"], label=label)

    for func in functions:
        nxt = func.get("NextFunctionID")
        if nxt and str(nxt).strip():
            edge_label = func["Outputs"][0] if func["Outputs"] else ""
            dot.edge(func["FunctionID"], nxt, label=edge_label, fontsize="9", fontname="Helvetica")

    svg_bytes = dot.pipe()
    out_path.write_bytes(svg_bytes)
    print(f"OK: {out_path}")


def main():
    functions = read_functions(DATA / "functions.xlsx")
    params_by_function = read_parameters(DATA / "parameters.xml")

    generate_matlab_script(functions, params_by_function, OUT / "build_avas_model.m")
    generate_preview_svg(functions, OUT / "architecture_preview.svg")

    print(f"\n{len(functions)} fonctions traitees.")
    print("Prochaine etape : ouvrir MATLAB, se placer dans ce dossier, et executer :")
    print("  run('output/build_avas_model.m')")


if __name__ == "__main__":
    main()
